"""Mengubah checklist peneliti (riset/checklist-peneliti.xlsx) menjadi berkas konten JSON di content/.

Isi pertanyaan, indikator, cara verifikasi, rekomendasi, dan kategori hasil disalin apa adanya dari sheet.
Kolom catatan internal (komentar antar-peneliti) tidak ikut disalin.
Jalankan ulang bila sheet berubah: python3 scripts/impor_xlsx.py
"""
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
wb = openpyxl.load_workbook(ROOT / "riset" / "checklist-peneliti.xlsx", data_only=True)

AREA = {
    "General Mobility": "mobilitas",
    "Dapur": "dapur",
    "Kamar Tidur": "kamar-tidur",
    "Kamar Mandi": "kamar-mandi",
}


def teks(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def rujukan(v):
    """Daftar rujukan: satu butir per baris, tanda '- ' di awal dibuang."""
    s = teks(v)
    if not s:
        return []
    butir = [re.sub(r"^\s*-\s*", "", b).strip() for b in re.split(r"\n+", s)]
    return [b for b in butir if b]


ws = wb["Pertanyaan Assesmen"]
per_area: dict[str, dict] = {}
sekarang = None
for r in ws.iter_rows(min_row=6, values_only=True):
    if r[0] in AREA and r[2] is None:
        sekarang = AREA[r[0]]
        per_area[sekarang] = {"fokus": None, "pertanyaan": []}
        continue
    if sekarang is None or not r[2]:
        continue
    m = re.match(r"^([A-Z]-\d+)\.\s*(.+)$", str(r[2]).strip())
    if not m:
        continue
    if r[0] and per_area[sekarang]["fokus"] is None:
        per_area[sekarang]["fokus"] = teks(r[0])
    per_area[sekarang]["pertanyaan"].append(
        {
            "id": m.group(1),
            "elemen": teks(r[1]),
            "pertanyaan": m.group(2).strip(),
            "indikator": teks(r[4]),
            "cara_verifikasi": teks(r[5]),
            "referensi": teks(r[6]),
            "hack": teks(r[8]),
            "modifikasi": teks(r[9]),
            "referensi_rekomendasi": rujukan(r[10]),
            "status_validasi": "draf",
        }
    )

(ROOT / "content" / "pertanyaan").mkdir(parents=True, exist_ok=True)
for area, isi in per_area.items():
    (ROOT / "content" / "pertanyaan" / f"{area}.json").write_text(
        json.dumps(isi, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(area, len(isi["pertanyaan"]), "pertanyaan")

# Kategori hasil (A-D) dan aturan penting
ws = wb["Kategori hasil asesmen"]
kategori, aturan = [], None
for r in ws.iter_rows(min_row=5, values_only=True):
    if r[0] in ("A", "B", "C", "D") and r[1]:
        kategori.append({"nilai": r[0], "status": teks(r[1]), "makna": teks(r[2]), "tindak_lanjut": teks(r[3])})
    elif r[1] and str(r[1]).startswith("Aturan penting"):
        aturan = teks(r[1])

# Tabel referensi
ws = wb["Tabel Referensi dan Petunjuk As"]
acuan, catatan = [], None
for r in ws.iter_rows(min_row=7, values_only=True):
    if r[0] and r[1] and not str(r[0]).startswith("Catatan"):
        acuan.append({"acuan": teks(r[0]), "penggunaan": teks(r[1])})
    elif r[0] and str(r[0]).startswith("Catatan"):
        catatan = teks(r[0])

(ROOT / "content" / "hasil.json").write_text(
    json.dumps({"kategori": kategori, "aturan_penting": aturan, "acuan": acuan, "catatan_acuan": catatan}, ensure_ascii=False, indent=2)
    + "\n",
    encoding="utf-8",
)
print("kategori", len(kategori), "acuan", len(acuan))
